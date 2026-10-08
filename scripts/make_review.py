"""由 pid_model.json 生成《P&ID 设计依据与连接关系表》xlsx（供用户确认连接关系后再绘图）。
用法: python make_review.py <pid_model.json> <输出.xlsx>
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

M = {}

FONT = Font(name='宋体', size=10)
BOLD = Font(name='宋体', size=10, bold=True)
TITLE = Font(name='宋体', size=14, bold=True)
THIN = Side(style='thin', color='000000')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill('solid', fgColor='D9D9D9')
WRAP = Alignment(wrap_text=True, vertical='center')
CENTER = Alignment(wrap_text=True, vertical='center', horizontal='center')

# 对象类别—采用符号—来源（阶段 2 配置时按此落实；尺寸为 A3 实际打印尺寸）
SYMBOLS = [
    ('主工艺管线（316L Φ6）', 'PipeLines / Primary', 'Primary Style', 'PIP', '线宽 0.5 mm', '管线号 HG/T 20519 格式'),
    ('驱动气管（PU Φ4）', 'PipeLines / Secondary', 'Secondary Style', 'PIP', '线宽 0.35 mm', ''),
    ('电信号线', 'SignalLines / Electrical', 'Electrical Style', 'PIP（ISA-5.1 表达）', '0.25 mm 虚线', ''),
    ('软件/通信链路', 'SignalLines / Software', 'Software Style', 'PIP（ISA-5.1 表达）', '0.25 mm', 'RS485/RS232/以太网'),
    ('气动信号', 'SignalLines / Pneumatic', 'Pneumatic Style', 'PIP（ISA-5.1 表达）', '0.25 mm', '先导阀→执行机构'),
    ('手动球阀', 'HandValves / BallValve', 'Ball Valve Style', 'PIP', '8×5 mm', ''),
    ('针阀', 'HandValves / NeedleValve', 'Needle Valve Style', 'PIP', '8×4 mm', '限流'),
    ('止回阀', 'HandValves / CheckValve', 'Check Valve Style', 'PIP', '8×4 mm', ''),
    ('气动隔膜切断阀（带先导电磁阀）', 'ControlValve 阀体 + Diaphragm 执行机构', 'Diaphragm Valve Style + 自定义“隔膜执行机构+先导电磁阀”', 'PIP 阀体 + ISA 执行机构（项目补充）', '阀体 8×4 mm', 'FC/FO 标注；先导电磁阀 XY 作执行机构一部分'),
    ('自力式减压阀', 'ControlValve 阀体 + PressureReducingRegulator', 'Control Valve Style + PressureReducingRegulator Style', 'PIP', '8×8 mm', '带压力表 PI'),
    ('管道过滤器', 'PipingSpecialtyItems / FilterPipe', 'Filter Pipe Style', 'PIP', '8×8 mm', ''),
    ('封堵帽', 'Fittings / Cap', 'Cap Style', 'PIP', '—', '备用口'),
    ('质量流量控制器 MFC', 'MassFlowController（新增类）', '自定义“MFC 本体”', '项目自定义符号', '14×6 mm', '本体位号 FIC-xxx'),
    ('流量开关', 'InlineInstruments / FlowSwitch', '自定义或 ISA 在线开关', 'ISA（项目补充）', '8×6 mm', 'FSL-201'),
    ('现场仪表（PT/TT/PI/AIT/TE）', 'GeneralInstrumentSymbols / FieldMountedDiscreteInstrument', 'Field Mounted Discrete Instrument Style', 'PIP（ISA-5.1 圆）', 'Ø9 mm', '字母代号+回路号'),
    ('PLC 功能（PI/TI/PAHH/TAHH 等）', 'GeneralInstrumentSymbols / PrimaryLocAccessiblePLC', 'Primary Loc Accessible PLC Style', 'PIP（ISA-5.1 方框内菱形）', '9×9 mm', ''),
    ('硬件联锁', 'FieldMountedDiscreteHardwareInterlock', 'Field Mounted Discrete Hardware Interlock Style', 'PIP（ISA-5.1 菱形 I）', '9×9 mm', '安全继电器 UZ-401'),
    ('气体探测器', 'GasDetector（新增类）', 'Field Mounted Discrete Instrument Style', 'PIP（ISA-5.1）', 'Ø9 mm', 'AIT-401~403'),
    ('管式炉', 'Furnaces / Furnace', '自定义“开启式管式炉”', '项目自定义符号', '约 60×25 mm', '内含夹具'),
    ('纽扣电池测试夹具', 'TestFixture（新增类）', '自定义“刚玉管夹具+纽扣电池”', '项目自定义符号', '约 50×10 mm', '阳极进/出管口'),
    ('冷凝收集杯', 'Tanks / Vessel', '自定义“小收集杯”', '项目自定义符号', '约 10×16 mm', '底部排液'),
    ('电子负载', 'ElectronicLoad（新增类）', '自定义“电子负载框”', '项目自定义符号', '约 30×16 mm', 'JL-301'),
    ('跨图连接器', 'Connectors / OffPageConnector', 'Off Page Connector Style', 'PIP', '—', '两端配对'),
    ('界区接口（甲方供/去向）', 'Connectors / OffPageConnector', 'Off Page Connector Style + 文字', 'PIP', '—', '注明甲方边界'),
]


def sheet(wb, name, headers, rows, widths, title=None):
    ws = wb.create_sheet(name)
    r0 = 1
    if title:
        ws.cell(1, 1, title).font = TITLE
        r0 = 3
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.border, c.alignment = BOLD, HEAD, BOX, CENTER
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font, c.border, c.alignment = FONT, BOX, WRAP
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{r0}:{r0}'
    return ws


def main(model, out):
    M.update(json.loads(Path(model).read_text(encoding='utf-8')))
    wb = Workbook()
    ws = wb.active
    ws.title = '说明'
    meta = M['meta']
    lines = [
        ('文件', 'P&ID 设计依据与连接关系表（供确认，确认后据此绘制 Plant 3D P&ID）'),
        ('项目', meta['project']),
        ('版次/日期', f"{meta['rev']} / {meta['date']}"),
        ('状态', meta['status']),
        ('设计依据', '\n'.join(meta['basis'])),
        ('依据类别代号', '；'.join(f'{k}={v}' for k, v in meta['source_codes'].items()) + '；R8=技术方案 R8，R1=成本核算 R1，B0=四件套 R0 baseline'),
        ('管线号规则', M['conventions']['line_number']),
        ('仪表位号规则', M['conventions']['tags']['instrument']),
        ('其他位号', '；'.join(f'{k}: {v}' for k, v in M['conventions']['tags'].items() if k != 'instrument')),
        ('管道等级', '；'.join(f'{k}: {v}' for k, v in M['conventions']['pipe_class'].items())),
        ('物料代号', '；'.join(f'{k}={v}' for k, v in M['conventions']['fluid_codes'].items())),
        ('图纸', '；'.join(f"{s['no']} {s['title']}" for s in M['sheets'])),
        ('暂定值', '注“暂定”的设定值在 FAT 前确定。'),
    ]
    ws.cell(1, 1, '青木子纽扣电池测试台 P&ID 设计依据与连接关系表').font = TITLE
    for i, (k, v) in enumerate(lines, 3):
        a, b = ws.cell(i, 1, k), ws.cell(i, 2, v)
        a.font, b.font = BOLD, FONT
        a.border = b.border = BOX
        a.alignment, b.alignment = CENTER, WRAP
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 16, 120

    sheet(wb, '界区接口', ['编号', '方向', '内容', '条件', '供方', '依据'],
          [[b['id'], '进' if b['dir'] == 'in' else '出', b['desc'], b['cond'], b['by'], b['src']] for b in M['boundary']],
          [10, 6, 40, 46, 10, 12])

    sheet(wb, '连接关系', ['管线号', '介质', '起点', '依次经过（流向）', '终点', '操作条件', '流量', '备注'],
          [[l['no'], M['conventions']['fluid_codes'].get(l['fluid'], l['fluid']), l['seq'][0], ' → '.join(l['seq'][1:-1]), l['seq'][-1],
            l.get('op', ''), l.get('flow', ''), l.get('note', '')] for l in M['lines']],
          [16, 14, 12, 60, 18, 22, 18, 30])

    sheet(wb, '三通与分支', ['编号', '所在管线', '分支', '说明'],
          [[t['id'], t['on'], t['branch'], t.get('desc', '')] for t in M['tees']], [10, 18, 18, 50])

    sheet(wb, '设备', ['位号', 'Plant 类', '符号', '名称', '规格', '供方', '依据'],
          [[e['tag'], e['class'], e['symbol'], e['name'], e['spec'], e['by'], e['src']] for e in M['equipment']],
          [10, 16, 18, 18, 70, 18, 14])

    sheet(wb, '阀门管件与在线仪表', ['位号', '类型', '说明', '规格', '故障位置', '依据'],
          [[x['tag'], x['type'], x['desc'], x.get('spec', ''), x.get('fail', ''), x['src']] for x in M['inline']],
          [10, 14, 34, 60, 8, 22])

    loc = {'field': '现场', 'field_integral': '现场（一体）', 'plc': 'PLC/HMI', 'vendor': '设备自带'}
    sheet(wb, '仪表与点位', ['位号', '位置', '功能', '量程', '设定值', 'I/O', '依据'],
          [[x['tag'], loc.get(x['loc'], x['loc']), x['desc'], x.get('range', ''), x.get('setpoint', x.get('alarms', '')), x.get('io', ''), x['src']]
           for x in M['instruments']],
          [10, 10, 40, 14, 34, 22, 24])

    typ = {'software': '软件/通信', 'electric': '电信号', 'pneumatic': '气动'}
    sheet(wb, '信号', ['起', '止', '类型', '说明'],
          [[s['from'], s['to'], typ.get(s['type'], s['type']), s.get('desc', '')] for s in M['signals']], [22, 30, 12, 60])

    acts = [('fuel_xv', '燃料切断阀 XV-101/103/104/105（含 XV-102）'), ('n2_bypass', 'N₂ 保护旁路 XV-201'),
            ('heat', '炉加热 XY-301'), ('load', '负载 XY-304'), ('alarm', '报警 XA-401'), ('permissive', '启动许可')]
    sheet(wb, '联锁因果表', ['编号', '原因', '检测', '设定值'] + [a[1] for a in acts] + ['复位', '依据'],
          [[i['id'], i['cause'], i['detect'], i['setpoint']] + [i['actions'].get(k, '') for k, _ in acts] + [i['reset'], i['src']]
           for i in M['interlocks']],
          [7, 26, 20, 18, 18, 16, 10, 8, 14, 18, 22, 18])

    sheet(wb, '运行工况', ['编号', '工况', '燃料切断阀', 'XV-102', 'N₂ 旁路 XV-201', 'N₂ 气源', 'MFC', '加热', '负载', '许可条件', '转出条件'],
          [[m['id'], m['name'], m['fuel_xv'], m['xv102'], m['bypass'], m['n2_source'], m['mfc'], m['heat'], m['load'], m['permissive'], m['exit']]
           for m in M['modes']],
          [6, 22, 18, 10, 18, 16, 24, 18, 18, 30, 26])

    sheet(wb, '符号对应表', ['对象类别', 'Plant 类', '采用符号（样式）', '来源', 'A3 打印尺寸', '注'], [list(r) for r in SYMBOLS],
          [26, 34, 40, 26, 14, 30], title='对象类别—采用符号—来源（计划，阶段 2 配置后以项目符号库为准）')

    sheet(wb, '待确认', ['编号', '事项', '原因', '建议'],
          [[q['id'], q['item'], q['why'], q['proposal']] for q in M['open_points']], [7, 40, 50, 40])

    wb.properties.creator = ''
    wb.properties.lastModifiedBy = ''
    wb.save(out)
    print('saved', out)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
