"""由 pid_model.json 生成《联锁因果表》xlsx 与《运行工况与联锁说明》docx（与图纸联锁矩阵、运行工况表同源）。
用法: python make_ce_modes.py <pid_model.json> <输出目录>
模型中可用的字段（缺省时用青木子项目的列定义）：
  meta.project / rev / date / status / basis_doc[] / ce_notes[]（xlsx 说明页补充条）/ safety_rules[]（docx 安全态与优先规则）/ pending_note
  conventions.action_columns  [[动作键, 表头, 简称], …]        联锁 actions 的列
  conventions.mode_columns    [[工况键, 表头], …]              modes 的列（permissive/exit 放最后）
  conventions.effect_columns  [动作列名, …] 与每条联锁的 effects[]：因果矩阵（推荐写明）
文字措辞：不写数据文件名、工具与过程；类别写“设计确认 / 设计依据 / 工程假设”。
依赖：openpyxl、python-docx；docx 版式用 docx-plain-technical-layout 技能的 plain_docx_builder.py。
"""
import copy
import importlib.util
import re
import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pidlib  # noqa: E402

PLAIN = Path(r'C:\Users\yujie\.codex\skills\docx-plain-technical-layout\scripts\plain_docx_builder.py')
_spec = importlib.util.spec_from_file_location('plain', PLAIN)
plain = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(plain)

M = pidlib.M
SUB = {'₂': '2', '₄': '4'}
SUP = {'²': '2', '³': '3'}
SCRIPT_RE = re.compile('([₂₄²³])')
ACT_DEFAULT = [['fuel_xv', '燃料切断阀 XV-101~105', '燃料阀'], ['n2_bypass', 'N₂ 旁路 XV-201', 'N₂ 旁路'],
               ['heat', '炉加热', '加热'], ['load', '电子负载', '负载'], ['alarm', '报警', '报警'], ['permissive', '运行许可', '许可']]
MODE_DEFAULT = [['fuel_xv', '燃料切断阀'], ['xv102', 'XV-102'], ['bypass', 'N₂ 旁路 XV-201'], ['n2_source', 'N₂ 气源 HV-102'],
                ['mfc', 'MFC'], ['heat', '加热'], ['load', '负载'], ['permissive', '许可条件'], ['exit', '转出条件']]


def cols():
    conv = M.get('conventions', {})
    return conv.get('action_columns') or ACT_DEFAULT, conv.get('mode_columns') or MODE_DEFAULT


def basis(src):
    """依据类别：设计确认 / 工程假设 / 设计依据（技术方案、成本等已有文件）。"""
    cats = []
    if '设计确认' in src:
        cats.append('设计确认')
    if '工程推荐' in src or '（E）' in src or '/E' in src or '设计复核' in src or '工程假设' in src:
        cats.append('工程假设')
    if re.search(r'\bR\d+\b', src):
        cats.append('设计依据')
    return '、'.join(cats) or '设计依据'


# ------------------------------------------------------------------ xlsx
FONT = Font(name='宋体', size=10)
BOLD = Font(name='宋体', size=10, bold=True)
TITLE = Font(name='宋体', size=14, bold=True)
THIN = Side(style='thin', color='000000')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill('solid', fgColor='D9D9D9')
WRAP = Alignment(wrap_text=True, vertical='center')
CENTER = Alignment(wrap_text=True, vertical='center', horizontal='center')


def rich(v, bold=False):
    """含下标/上标字符的文字转为 Excel 富文本（真实上下标格式）。"""
    if not isinstance(v, str) or not SCRIPT_RE.search(v):
        return v
    parts = []
    for tok in SCRIPT_RE.split(v):
        if not tok:
            continue
        if tok in SUB:
            parts.append(TextBlock(InlineFont(rFont='宋体', sz=10, b=bold, vertAlign='subscript'), SUB[tok]))
        elif tok in SUP:
            parts.append(TextBlock(InlineFont(rFont='宋体', sz=10, b=bold, vertAlign='superscript'), SUP[tok]))
        else:
            parts.append(TextBlock(InlineFont(rFont='宋体', sz=10, b=bold), tok))
    return CellRichText(parts)


def lines(v, width):
    """单元格换行后的行数估算：全角字符按 2 个字符宽，列宽留 1 个字符余量（Excel 自适应对中文估计偏小）。"""
    if not v:
        return 1
    n = 0
    for seg in str(v).split('\n'):
        w = sum(2 if ord(ch) > 0x2E80 else 1.05 for ch in seg)
        n += max(1, -(-int(w) // max(1, int(width) - 1)))
    return n


def sheet(wb, name, title, headers, rows, widths, center_cols=()):
    ws = wb.create_sheet(name)
    ws.cell(1, 1, title).font = TITLE
    r0 = 3
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, rich(h, True))
        c.font, c.fill, c.border, c.alignment = BOLD, HEAD, BOX, CENTER
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, rich(v))
            c.font, c.border = FONT, BOX
            c.alignment = CENTER if j in center_cols else WRAP
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, row in enumerate([headers] + rows, r0):
        ws.row_dimensions[i].height = max(lines(v, w) for v, w in zip(row, widths)) * 13.5 + 5
    ws.freeze_panes = ws.cell(r0 + 1, 3)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{r0}:{r0}'
    ws.oddFooter.center.text = '第 &P 页 / 共 &N 页'
    return ws


def build_xlsx(out):
    act, mode = cols()
    meta = M['meta']
    wb = Workbook()
    ws = wb.active
    ws.title = '说明'
    notes = [(f"{meta['project']} 联锁因果表（C&E）与运行工况", TITLE),
             (f"版次：{meta.get('rev', 'R0')}    日期：{meta.get('date', '')}    状态：{meta.get('status', '')}", FONT),
             ('设计依据：' + '；'.join(meta.get('basis_doc', [])) + '。', FONT),
             ('依据类别：设计确认＝已确认的设计决定；设计依据＝技术方案、成本核算等已有文件中的内容；工程假设＝本阶段提出、FAT 前验证的内容。', FONT)]
    notes += [(t, FONT) for t in meta.get('ce_notes', [])]
    notes.append(('注“暂定/暂值”的设定值在 FAT 前确定。', FONT))
    for r, (t, f) in enumerate(notes, 1):
        c = ws.cell(r, 1, rich(t, f is TITLE))
        c.font, c.alignment = f, Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 150
    for r in range(2, len(notes) + 1):
        ws.row_dimensions[r].height = 30
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3

    rows = [[i['id'], i['cause'], i.get('detect', ''), i.get('setpoint', '')] + [i['actions'].get(k, '') for k, _, _ in act]
            + [i.get('reset', ''), basis(i.get('src', '')), i.get('src', '')] for i in M['interlocks']]
    sheet(wb, '联锁因果表', '联锁因果表（原因 → 动作；空白＝不动作）',
          ['编号', '原因', '检测/执行', '设定值'] + [h for _, h, _ in act] + ['复位', '依据类别', '依据出处'],
          rows, [8, 30, 26, 16] + [22] * len(act) + [28, 18, 34], center_cols=(1,))

    ecols = list(pidlib.ce_matrix(M['interlocks'][0]).keys())
    rows = [[i['id'], i['cause']] + list(pidlib.ce_matrix(i).values()) for i in M['interlocks']]
    ws = sheet(wb, '因果矩阵', '因果矩阵（X＝该原因触发该动作）', ['编号', '原因'] + ecols, rows,
               [8, 34] + [11] * len(ecols), center_cols=tuple(range(3, 3 + len(ecols))) + (1,))
    ws.row_dimensions[3].height = 45

    rows = [[m['id'], m['name']] + [m.get(k, '') for k, _ in mode] for m in M['modes']]
    sheet(wb, '运行工况', '运行工况表（各工况阀位、流量、加热、负载、许可与转出条件）',
          ['编号', '工况'] + [h for _, h in mode], rows, [7, 22] + [22] * len(mode), center_cols=(1,))

    rows = [[q['id'], q['item'], q.get('why', ''), q.get('proposal', ''), q.get('resolution', '待定')] for q in M.get('open_points', [])]
    sheet(wb, '待确认与决定', '待确认事项与决定记录', ['编号', '事项', '原因', '建议', '决定'], rows,
          [8, 40, 46, 46, 46], center_cols=(1,))
    wb.properties.creator = ''
    wb.properties.lastModifiedBy = ''
    wb.save(out)


# ------------------------------------------------------------------ docx
def split_scripts(doc):
    """把 run 中的 ₂ ₄ ² ³ 拆成真实下标/上标 run。"""
    from docx.text.run import Run

    def paras(container):
        for p in container.paragraphs:
            yield p
        for t in container.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from paras(cell)
    for p in paras(doc):
        for run in list(p.runs):
            if not SCRIPT_RE.search(run.text):
                continue
            toks = [t for t in SCRIPT_RE.split(run.text) if t]
            run.text = ''
            anchor = run._r
            for tok in toks:
                r = copy.deepcopy(run._r)
                for t in r.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
                    r.remove(t)
                anchor.addnext(r)
                anchor = r
                nr = Run(r, p)
                if tok in SUB:
                    nr.text, nr.font.subscript = SUB[tok], True
                elif tok in SUP:
                    nr.text, nr.font.superscript = SUP[tok], True
                else:
                    nr.text = tok
            run._r.getparent().remove(run._r)


def dtable(doc, width, n, cap, headers, rows, weights):
    plain.caption(doc, f'表{n}  {cap}')
    t = plain.make_table(doc, width, headers, rows, weights)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(plain.element('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(out):
    act, mode = cols()
    meta = M['meta']
    name = meta['project']
    doc, width, num = plain.setup_document()
    doc.core_properties.title = f'{name}运行工况与联锁说明'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''
    t = doc.add_paragraph(style='Title')
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    plain.font(t.add_run(f'{name}运行工况与联锁说明'), 20, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    plain.font(p.add_run(f"{meta.get('rev', 'R0')}    {meta.get('date', '')}"), 10.5)
    hdr = doc.sections[0].header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    plain.font(hdr.add_run(f'{name}  运行工况与联锁说明'), 9)

    plain.heading(doc, '范围与依据', 1, num)
    plain.body(doc, meta.get('scope_text', f'本说明给出{name}的联锁因果关系和运行工况；P&ID 中的联锁矩阵与运行工况表为本说明的简表。'))
    plain.body(doc, '设计依据：' + '；'.join(meta.get('basis_doc', [])) + '。')
    plain.body(doc, '依据类别分三类标注：“设计确认”为已确认的设计决定；“设计依据”为技术方案、成本核算等已有文件中的内容；'
                    '“工程假设”为本阶段提出，须在 FAT 前验证或冻结的内容。标“暂定”或“暂值”的设定值均属后者。')
    if meta.get('safety_rules'):
        plain.heading(doc, '安全态与优先规则', 1, num)
        for s in meta['safety_rules']:
            plain.body(doc, s)

    plain.heading(doc, '联锁因果表', 1, num)
    plain.body(doc, '表1 列出各联锁的原因、检测方式、设定值、动作和复位要求；动作栏空白表示该联锁不触发对应动作。')
    rows = []
    for i in M['interlocks']:
        a = i['actions']
        txt = '；'.join(f'{s}：{a[k]}' for k, _, s in act if a.get(k))
        rows.append([i['id'], f"{i['cause']}\n检测：{i.get('detect', '')}", i.get('setpoint', ''), txt, i.get('reset', ''), basis(i.get('src', ''))])
    dtable(doc, width, 1, '联锁因果表', ['编号', '原因与检测', '设定值', '动作', '复位', '依据类别'], rows, [7, 24, 11, 33, 14, 11])

    plain.heading(doc, '运行工况', 1, num)
    plain.body(doc, '表2 列出各运行工况的阀位、流量、加热、负载、许可与转出条件。工况转换须经 HMI 确认。')
    state = [(k, h) for k, h in mode if k not in ('permissive', 'exit')]
    rows = [[m['id'], m['name'], '\n'.join(f'{h}：{m.get(k, "")}' for k, h in state), m.get('permissive', ''), m.get('exit', '')]
            for m in M['modes']]
    dtable(doc, width, 2, '运行工况表', ['编号', '工况', '状态', '许可条件', '转出条件'], rows, [7, 14, 47, 16, 16])

    if M.get('open_points'):
        plain.heading(doc, '待确认事项与决定记录', 1, num)
        rows = [[q['id'], q['item'], q.get('resolution', '待定（建议：' + q.get('proposal', '') + '）')] for q in M['open_points']]
        dtable(doc, width, 3, '待确认事项与决定记录', ['编号', '事项', '决定或建议'], rows, [8, 42, 50])
    if meta.get('pending_note'):
        plain.heading(doc, '有待确认的设定与接口', 1, num)
        plain.body(doc, meta['pending_note'])
    for node in list(doc.element.xpath('//w:pBdr')) + list(doc.styles['Title'].element.xpath('.//w:pBdr')):
        node.getparent().remove(node)
    doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    split_scripts(doc)
    doc.save(out)


if __name__ == '__main__':
    pidlib.use_model(sys.argv[1])
    od = Path(sys.argv[2])
    od.mkdir(parents=True, exist_ok=True)
    build_xlsx(od / '联锁因果表.xlsx')
    build_docx(od / '运行工况与联锁说明.docx')
    print('ok')
