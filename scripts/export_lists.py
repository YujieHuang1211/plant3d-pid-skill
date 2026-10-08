"""由 pid_model.json 导出设备表、管线表、仪表索引、阀门及管道附件表、界区接口表，并与成本 R1 询价表、控制点位表对账。
用法: python export_lists.py <pid_model.json> <输出.xlsx> [成本表.xlsx|-] [ProcessPower.dcf]
  成本表给“-”则不做询价表/控制点位表对账；一个询价行对应多个图面对象时，在模型 conventions.wbs_per_unit 中写 {"02.01": [2, "说明"]}
  - 给出 dcf 时，再核对 Plant 项目库中的位号与模型一致（双向）
  - 询价表按 WBS 对数量：模型对象的 src 中 “R1 xx.yy” 计数 vs 询价表数量（按件计量的单位才对数）
  - 控制点位表：模型 io 字段引用的点位 vs 点位表，双向列出差异
对账结果另存 <输出>.recon.json；只读 R1，不改上游文件。
"""
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pidlib  # noqa: E402
from make_ce_modes import rich, sheet  # noqa: E402
from verify_topology import load_db, tag_map  # noqa: E402

M = pidlib.M
TYPE_CN = {'ball_valve': '手动球阀', 'needle_valve': '针阀', 'check_valve': '止回阀', 'filter': '过滤器', 'regulator': '减压阀',
           'on_off_valve': '气动隔膜切断阀', 'mfc': '质量流量控制器', 'cap': '封堵帽', 'flow_switch': '流量开关'}
COUNT_UNITS = {'台', '只', '套', '支', '块'}
# R1 询价表中一个计价单位在 P&ID 上对应多个图面对象（依据 R1 该行“成套配置/设备说明”原文）
PER_UNIT = {
    '02.01': (2, '每套＝减压阀 PCV + 压力表 PI'),
    '02.07': (3, '成套配置含隔膜阀（XV-201）及驱动气减压阀与压力表 1 套（PCV-202、PI-202）'),
    '02.08': (2, '针阀 HV-202 + 流量开关 FSL-201'),
    '02.15': (2, '收集杯 V-301 + 排液阀 HV-301'),
    '03.01': (2, '管式炉 F-301 + 炉自带控制热偶 TE-301'),
    '03.06': (2, 'Pt100 TE-303 + 变送器 TT-303'),
    '05.12': (3, '急停 HS-401 + 复位 HS-402 + 总声光 XA-401'),
}
WBS_RE = re.compile(r'R1 (\d\d\.\d\d)')
IO_RE = re.compile(r'\b(DI\d\d|DO\d\d|AI\d\d|TC\d\d|COM\d\d|NET\d\d)\b')


def wbs_of(src):
    """src 中第一个 R1 WBS（同一对象只计一次）；“R1 02.01/B0 D09”→02.01。"""
    m = WBS_RE.search(src or '')
    return m.group(1) if m else None


def main(model, out, r1='-', dcf=None):
    pidlib.use_model(model)
    PER_UNIT.update({k: tuple(v) for k, v in M.get('conventions', {}).get('wbs_per_unit', {}).items()})
    out = Path(out)
    wb = Workbook()
    wb.remove(wb.active)

    eq = M['equipment']
    sheet(wb, '设备表', '设备表', ['位号', '名称', '规格/要求', '供货', 'Plant 类', '依据'],
          [[e['tag'], e['name'], e['spec'], e.get('by', ''), e['class'], e.get('src', '')] for e in eq], [10, 22, 70, 22, 16, 18], (1,))

    rows = []
    for ln in M['lines']:
        parts = ln['no'].split('-')
        rows.append([ln['no'], parts[0], parts[2], parts[3], ln['seq'][0], ln['seq'][-1].split(':')[0], ' → '.join(s.split(':')[0] for s in ln['seq']),
                     ln.get('op', ''), ln.get('flow', '')])
    sheet(wb, '管线表', '管线表（管线号：物料代号-主项序号-外径mm-管道等级，HG/T 20519）',
          ['管线号', '物料', '外径 mm', '等级', '起点', '终点', '途经', '操作条件', '流量'], rows, [18, 7, 8, 6, 12, 16, 70, 30, 16], (1,))

    sheet(wb, '仪表索引', '仪表索引（ISA-5.1 字母代号 + 回路号）',
          ['位号', '说明', '安装位置', '量程', '设定值/报警', 'I/O 点位', '所在对象', '依据'],
          [[i['tag'], i['desc'], i['loc'], i.get('range', ''), '；'.join(x for x in [i.get('setpoint', ''), i.get('alarms', '') if isinstance(i.get('alarms'), str) else '；'.join(i.get('alarms', []) or [])] if x),
            i.get('io', ''), i.get('on', ''), i.get('src', '')] for i in M['instruments']],
          [11, 40, 10, 16, 24, 22, 12, 26], (1,))

    sheet(wb, '阀门及管道附件表', '阀门及管道附件表',
          ['位号', '类型', '说明', '规格', '依据'],
          [[v['tag'], TYPE_CN.get(v['type'], v['type']), v.get('desc', ''), v.get('spec', ''), v.get('src', '')] for v in M['inline']],
          [10, 16, 40, 50, 22], (1,))

    sheet(wb, '界区接口表', '界区接口表（甲方/我方分界）', ['编号', '方向', '说明', '条件', '责任方', '依据'],
          [[b['id'], {'in': '进', 'out': '出'}.get(b.get('dir'), b.get('dir', '')), b['desc'], b.get('cond', ''), b.get('by', ''), b.get('src', '')] for b in M['boundary']],
          [10, 6, 50, 26, 12, 18], (1,))

    rec_q, rec_io = [], []
    if r1 and r1 != '-':
        rec_q, rec_io = reconcile(wb, eq, r1)

    # ---------- 核对：Plant 项目库位号
    rec_db = None
    if dcf:
        rec_db = check_db(wb, eq, dcf)
    finish(wb, out, rec_q, rec_io, rec_db)


def reconcile(wb, eq, r1):
    # ---------- 对账：询价表
    r1wb = openpyxl.load_workbook(r1, data_only=True, read_only=True)
    q = {}
    for r in r1wb['询价表'].iter_rows(min_row=2, values_only=True):
        if r[0] and re.fullmatch(r'\d\d\.\d\d', str(r[0])):
            q[str(r[0])] = (r[2], r[5], r[6])
    used = defaultdict(list)
    for obj in eq + M['inline'] + M['instruments']:
        w = wbs_of(obj.get('src'))
        if w:
            used[w].append(obj['tag'])
    rec_q, rows = [], []
    for w, (name, qty, unit) in q.items():
        tags = used.get(w, [])
        if unit in COUNT_UNITS and tags:
            per, why = PER_UNIT.get(w, (1, ''))
            exp = qty * per
            st = ('一致' if len(tags) == exp else f'不一致：图上 {len(tags)}，应为 {exp}') + (f'（{why}）' if why else '')
        elif tags:
            st = '按批/长度计量，不对个数'
        elif w[:2] in ('01', '02', '03', '04', '05'):
            st = '图上未单列（附件/材料/系统内件，人工确认）'
        else:
            st = '非 P&ID 对象（框架、人工、物流）'
        rows.append([w, name, qty, unit, len(tags), ' '.join(tags), st])
        rec_q.append({'wbs': w, 'name': name, 'qty': qty, 'unit': unit, 'tags': tags, 'status': st})
    for w in sorted(set(used) - set(q)):
        rows.append([w, '（询价表无此 WBS）', '', '', len(used[w]), ' '.join(used[w]), '模型引用了不存在的 WBS'])
        rec_q.append({'wbs': w, 'tags': used[w], 'status': '模型引用了不存在的 WBS'})
    sheet(wb, '对账_询价表', '与成本 R1 询价表对账（按 WBS；只读 R1）',
          ['WBS', '询价表名称', '数量', '单位', '图上对象数', '图上位号', '结论'], rows, [8, 26, 6, 6, 8, 60, 34], (1, 3, 4, 5))

    # ---------- 对账：控制点位表
    pts = {}
    for r in r1wb['控制点位表'].iter_rows(min_row=2, values_only=True):
        if r[0] and IO_RE.fullmatch(str(r[0])):
            pts[str(r[0])] = (r[2], r[3], r[7])
    ref = defaultdict(list)
    for i in M['instruments']:
        for p in IO_RE.findall(i.get('io', '')):
            ref[p].append(i['tag'])
    rec_io, rows = [], []
    for p, (fn, dev, note) in pts.items():
        tags = ref.get(p, [])
        note = note or ''
        if tags and ('释放' in note or '取消' in (fn or '')):
            st = '点位表为备用/取消，P&ID 已启用 → 上游待更新'
        elif tags:
            st = '一致'
        elif '释放' in note or '备用' in note or '取消' in (fn or ''):
            st = '备用（图上不画）'
        else:
            st = '图上无仪表引用此点位'
        rows.append([p, fn, dev, note, ' '.join(tags), st])
        rec_io.append({'point': p, 'function': fn, 'tags': tags, 'status': st})
    for p in sorted(set(ref) - set(pts)):
        rows.append([p, '（点位表无此点）', '', '', ' '.join(ref[p]), '模型引用了点位表中没有的点'])
        rec_io.append({'point': p, 'tags': ref[p], 'status': '模型引用了点位表中没有的点'})
    sheet(wb, '对账_控制点位表', '与成本 R1 控制点位表对账（只读 R1）',
          ['点位', '点位表功能', '设备/对象', '点位表说明', '图上位号', '结论'], rows, [8, 24, 22, 40, 24, 34], (1,))
    return rec_q, rec_io


def check_db(wb, eq, dcf):
    c = load_db(dcf)
    db_tags = {t for _, t in tag_map(c).values() if '?' not in t}
    want = {e['tag'] for e in eq} | {v['tag'] for v in M['inline']} | {i['tag'] for i in M['instruments']} | {ln['no'] for ln in M['lines']}
    rec_db = {'missing_in_plant': sorted(want - db_tags), 'extra_in_plant': sorted(db_tags - want)}
    rows = [[t, '模型有、Plant 库无'] for t in rec_db['missing_in_plant']] + [[t, 'Plant 库有、模型无（联锁/连接器/管嘴等图面对象可不在模型清单）'] for t in rec_db['extra_in_plant']]
    sheet(wb, '核对_Plant库位号', 'Plant 项目库位号与模型核对', ['位号', '差异'], rows or [['—', '无差异']], [24, 60], (1,))
    return rec_db


def finish(wb, out, rec_q, rec_io, rec_db):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = rich(cell.value, cell.font.b)
    wb.properties.creator = ''
    wb.properties.lastModifiedBy = ''
    wb.save(out)
    rec = {'quote': rec_q, 'io': rec_io, 'plant_db': rec_db}
    out.with_suffix('.recon.json').write_text(json.dumps(rec, ensure_ascii=False, indent=1), encoding='utf-8')
    bad_q = [r for r in rec_q if r['status'].startswith(('不一致', '模型引用'))]
    bad_io = [r for r in rec_io if r['status'] not in ('一致', '备用（图上不画）')]
    print('quote issues', len(bad_q), 'io issues', len(bad_io), 'plant', None if rec_db is None else (len(rec_db['missing_in_plant']), len(rec_db['extra_in_plant'])))


if __name__ == '__main__':
    main(*sys.argv[1:5])
