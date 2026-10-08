"""原生 Office 往返与渲染核查：《运行工况与联锁说明》docx、《联锁因果表》xlsx。
用法: python render_ce_modes.py <交付目录> <QA输出目录>
Word：打开→更新域→保存→关闭→只读重开→导出 PDF，文字与表格内容不变；
Excel：打开→CalculateFullRebuild→保存→关闭→重开→导出 PDF，单元格值不变、无错误值。
"""
import hashlib
import json
import os
import shutil
import sys
from pathlib import Path

import fitz
import openpyxl
import win32com.client as win32
from docx import Document


def docx_content(f):
    d = Document(f)
    return ([p.text for p in d.paragraphs], [[[c.text for c in r.cells] for r in t.rows] for t in d.tables])


def xlsx_values(f):
    wb = openpyxl.load_workbook(f)
    return {ws.title: [[str(c.value) if c.value is not None else None for c in r] for r in ws.iter_rows()] for ws in wb}


def pages(pdf, qa, stem):
    doc = fitz.open(pdf)
    out = []
    for i, p in enumerate(doc):
        png = qa / f'{stem}-p{i + 1:02}.png'
        p.get_pixmap(matrix=fitz.Matrix(1.6, 1.6)).save(png)
        out.append({'page': i + 1, 'png': png.name, 'chars': len(p.get_text())})
    return out


def main(od, qa):
    od, qa = Path(od), Path(qa)
    qa.mkdir(parents=True, exist_ok=True)
    local = Path(os.environ['TEMP']) / 'qingmuzi_pid_ce'
    local.mkdir(exist_ok=True)
    rep = {}

    src = od / '运行工况与联锁说明.docx'
    tgt = local / src.name
    shutil.copy2(src, tgt)
    before = docx_content(src)
    w = win32.DispatchEx('Word.Application')
    w.Visible, w.DisplayAlerts = False, 0
    try:
        d = w.Documents.Open(str(tgt), ReadOnly=False, AddToRecentFiles=False, OpenAndRepair=False)
        d.Repaginate()
        d.Fields.Update()
        d.Save()
        d.Close(False)
        d = w.Documents.Open(str(tgt), ReadOnly=True, AddToRecentFiles=False, OpenAndRepair=False)
        pdf = qa / (src.stem + '.pdf')
        d.ExportAsFixedFormat(str(pdf), 17)
        d.Close(False)
    finally:
        w.Quit()
    same = before == docx_content(tgt)
    if same:
        shutil.copy2(tgt, src)
    rep['docx'] = {'file': src.name, 'native_open_save_reopen': True, 'content_unchanged': same,
                   'sha256': hashlib.sha256(src.read_bytes()).hexdigest(), 'pages': pages(pdf, qa, src.stem)}

    src = od / '联锁因果表.xlsx'
    tgt = local / src.name
    shutil.copy2(src, tgt)
    before = xlsx_values(src)
    x = win32.DispatchEx('Excel.Application')
    x.Visible, x.DisplayAlerts = False, False
    errs = []
    try:
        wb = x.Workbooks.Open(str(tgt))
        x.CalculateFullRebuild()
        wb.Save()
        wb.Close(False)
        wb = x.Workbooks.Open(str(tgt), ReadOnly=True)
        for ws in wb.Worksheets:
            ur = ws.UsedRange
            vals = ur.Value
            for r, row in enumerate(vals if isinstance(vals, tuple) else ((vals,),)):
                for c, v in enumerate(row):
                    if isinstance(v, int) and v < 0:      # COM 以负整数返回 #N/A 等错误值
                        errs.append(f'{ws.Name}!R{r + 1}C{c + 1}')
                    if isinstance(v, str) and v.startswith('#'):
                        errs.append(f'{ws.Name}!R{r + 1}C{c + 1} {v}')
        pdf = qa / (src.stem + '.pdf')
        wb.ExportAsFixedFormat(0, str(pdf))
        wb.Close(False)
    finally:
        x.Quit()
    same = before == xlsx_values(tgt)
    if same:
        shutil.copy2(tgt, src)
    rep['xlsx'] = {'file': src.name, 'native_open_recalc_save_reopen': True, 'content_unchanged': same, 'error_cells': errs,
                   'sha256': hashlib.sha256(src.read_bytes()).hexdigest(), 'pages': pages(pdf, qa, src.stem)}
    rep['boundary'] = '文档往返与渲染检查，不等于 HAZOP、安全验收或实物联锁测试。'
    (qa / '联锁因果表与工况说明_渲染核查.json').write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding='utf-8')
    print('docx', rep['docx']['content_unchanged'], len(rep['docx']['pages']), 'xlsx', rep['xlsx']['content_unchanged'],
          len(rep['xlsx']['pages']), 'errs', len(errs))


if __name__ == '__main__':
    main(*sys.argv[1:3])
